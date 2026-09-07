from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

root = Path('D:/Temp/ClawAssistant')
doc = Document(root / '知行作品报告书_参考稿融合版.docx')

def replace(p, text):
    p.clear()
    r = p.add_run(text)
    r.font.name = '宋体'
    r.font.size = Pt(10.5)
    r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体')

replace(doc.tables[0].rows[1].cells[1].paragraphs[0], '主题1：大模型驱动的智能交互与对话系统')
replace(doc.tables[0].rows[3].cells[1].paragraphs[0], '知行是一套面向校园与日常生活场景的多用户 Web 智能体交互系统。系统以大模型对话为统一入口，结合 ReAct 决策与工具调用，根据用户意图选择天气、课表、考试、旅行、地图、文件和目标跟进等技能，通过 SSE 返回流式回答与任务状态。多轮上下文、长期记忆和任务续接支撑连续交互；图片、音频和文件处理扩展任务输入与输出。项目以智能交互与对话系统为主体，以知识记忆和多模态能力为支撑，将需求理解、工具执行、结果反馈和后续提醒接入同一工作台。')

for p in doc.paragraphs:
    t = p.text
    if t.startswith('知行是一款运行在浏览器中的个人智能助理。'):
        replace(p, '知行选择“主题1：大模型驱动的智能交互与对话系统”。作品面向校园与日常生活，以浏览器中的多用户对话工作台为入口：用户描述目标，系统结合上下文识别意图、补齐条件并选择技能与工具，再反馈执行结果。设计重点是让对话能够推进实际任务，而不只是生成建议。知识记忆用于保留背景与偏好，多模态能力用于接收资料和交付产物，两者共同支撑连续交互。')
    elif t.startswith('WebChatController 从认证会话取得用户身份'):
        replace(p, '对话主链路为 WebChatController → ChatApplicationService → ReActAgentExecutor → SkillRouter / ExecutionLoop / Tools。入口校验用户与会话归属、建立 runId 并解析附件；执行器组织技能路由、上下文和动态工具集，再由 ReAct 循环依据工具结果继续决策或结束回复。')
    elif t.startswith('对用户而言，这条链路可以概括为'):
        replace(p, '这条链路对应“理解需求—规划执行—工具协同—结果反馈”。用户补充条件或调整要求时，系统结合会话状态继续处理；信息不足时追问，高风险操作等待确认。SSE 将回答和工具状态推送到前端，用户可以核对调用过程、失败位置和生成产物。交互因此既包含自然语言交流，也包含对执行过程的确认与修正。')
    elif t.startswith('知行把 Agent 能力落成用户可以观察的步骤'):
        replace(p, '交互层将工具调用状态、待确认操作和业务结果呈现给用户，使用户能够参与执行过程。可核对的工具记录是判断任务完成的依据，不把模型生成的解释本身当作执行成功的证明。')
    elif t.startswith('当前版本已经连通任务执行、多模态资料处理'):
        replace(p, '知行以大模型驱动的智能交互系统为主体，以知识记忆和多模态为支撑。后续优先验证多轮续接、工具执行、确认交互和主动提醒的完成质量，再扩展场景。')
    elif t.startswith('演示一：登录知行'):
        replace(p, '演示一：上传课表，补充学期与提醒条件，检查追问、解析和查询是否衔接。演示二：提出天气与路线需求，再修改目的地，检查技能切换、上下文续接及工具结果。演示三：保存个人偏好，在记忆页核对后开启新一轮对话，检查召回是否相关。')
    # Avoid a mostly empty overflow page while preserving the original font and line spacing.
    if p.text.startswith('六 验证设计与演示安排'):
        p.paragraph_format.page_break_before = False

doc.core_properties.author = ''
doc.core_properties.last_modified_by = ''
doc.save(root / '知行作品报告书_主题1修订版.docx')
print('Theme 1 revision saved')
